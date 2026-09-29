"""Презентация в шаблоне организатора (ЛЦТ2026): обязательные слайды + блок решения.

    python make_presentation.py <шаблон.pptx> <выход.pptx>

Данные о команде — заглушки в [квадратных скобках]. Цифры берутся из out/metrics_final.json.
Картинки: out/pres/*.png (make_presentation_assets), out/slides/*.png (make_slides.py).
"""

import copy
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = Path(__file__).parent
TEMPLATE, OUT = sys.argv[1], sys.argv[2]
U = 91440  # «единица» = 0.1 дюйма (координаты макетов шаблона)
PLUM, PINK, DARK, GREY = RGBColor(0x31, 0x0F, 0x53), RGBColor(0xFF, 0x00, 0x53), RGBColor(0x1C, 0x1D, 0x22), RGBColor(0x55, 0x55, 0x60)
FONT = "Montserrat"

m = json.load(open(ROOT / "out/metrics_final.json", encoding="utf-8"))
M = lambda k, f: m[k][f]
MACRO = m["macro-F1 по типам нарушений"]
QP = m["quality_prob ROC-AUC: все снимки"]["auc"]
ALL = "все снимки: качество (есть нарушение)"

prs = Presentation(TEMPLATE)
src = list(prs.slides)  # исходные слайды шаблона (1-based в комментариях)
final = []


def dup(n):
    """Копия слайда шаблона n (1-based): фигуры и связанные картинки; диаграммы не копируются."""
    s = src[n - 1]
    new = prs.slides.add_slide(s.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rel in s.part.rels.values():
        if "image" in rel.reltype:
            rid_map[rel.rId] = new.part.relate_to(rel._target, rel.reltype)
    for sh in s.shapes:
        if sh.has_chart if hasattr(sh, "has_chart") else False:
            continue
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in ("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
                         "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"):
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        new.shapes._spTree.append(el)
    final.append(new)
    return new


def keep(n):
    final.append(src[n - 1])
    return src[n - 1]


def shape(slide, name, nth=0):
    found = [sh for sh in slide.shapes if sh.name == name]
    return found[nth] if len(found) > nth else None


def put(sh, lines, size=None, bold_first=False, color=None):
    """Заменить текст, сохранив оформление первого абзаца/пробега фигуры."""
    if sh is None:
        return
    lines = [lines] if isinstance(lines, str) else lines
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr = copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for child in list(p0._p):  # пробеги, переносы строк и поля — всё, кроме свойств абзаца
        if not child.tag.endswith("}pPr"):
            p0._p.remove(child)
    for k, line in enumerate(lines):
        p = p0 if k == 0 else tf.add_paragraph()
        if k and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        run = p.add_run()
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        run.text = line
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        elif rpr is None:  # пустой плейсхолдер шаблона: у него светлый текст — на белых карточках не виден
            run.font.color.rgb = DARK
        if bold_first and k == 0:
            run.font.bold = True


def text(slide, x, y, w, h, lines, size=14, color=DARK, bold_first=False, align=None):
    tb = slide.shapes.add_textbox(Emu(int(x * U)), Emu(int(y * U)), Emu(int(w * U)), Emu(int(h * U)))
    tf = tb.text_frame
    tf.word_wrap = True
    for k, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
        r.font.bold = bold_first and k == 0
        p.space_after = Pt(4)
        if align is not None:
            p.alignment = align
    return tb


def pic(slide, path, x, y, w, h):
    """Картинка, вписанная в прямоугольник (x, y, w, h) с сохранением пропорций, по центру."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    slide.shapes.add_picture(str(path), Emu(int((x + (w - pw) / 2) * U)), Emu(int((y + (h - ph) / 2) * U)),
                             Emu(int(pw * U)), Emu(int(ph * U)))


def drop(slide, *names):
    for n in names:
        sh = shape(slide, n)
        if sh is not None:
            sh._element.getparent().remove(sh._element)


def title(slide, t, name="Заголовок 13"):
    put(shape(slide, name), t, color=RGBColor(255, 255, 255))
    for sh in slide.shapes:  # плашка под заголовком — скруглённый прямоугольник вверху слайда
        if sh.shape_type == 1 and sh.top is not None and sh.top < 5 * U and sh.height < 8 * U:
            sh.width = Emu(int(max(sh.width / U, min(118, 6 + len(t) * 2.3)) * U))


# ---------------------------------------------------------------- обязательный блок
s = keep(7)
put(shape(s, "Заголовок 2"), "DXA-QC: ИИ-контроль качества денситометрии", color=RGBColor(255, 255, 255))
put(shape(s, "Текст 4"), "[Название команды] · Сервис ИИ по оценке качества исследований плотности костей", color=RGBColor(255, 255, 255))

s = keep(8)
put(shape(s, "Заголовок 6"), "[Название команды]", color=PLUM)
put(shape(s, "Текст 8", 2), ["Капитан: [ФИО, специальность]", "Кол-во участников: [N] человек",
                             "Краткое описание: [как образовалась команда, место учёбы/работы]",
                             "Город и регион: [город, регион]"])
put(shape(s, "Текст 8", 5), "Сервис проверяет DICOM-исследование DXA по критериям ТЗ и объясняет каждое "
                            "нарушение числом и порогом: наклон оси 7° при допуске 5°, 2 см ниже малого вертела при норме 3.")
put(shape(s, "Текст 8", 1), "Меряем то, что ТЗ задаёт числом, и учим только то, что формулой не описать: "
                            "решение проверяемо врачом и не переобучается на 6–36 примерах нарушений.")

s = keep(9)
title(s, "Команда", "Заголовок 6")
for sh in [x for x in s.shapes if x.has_text_frame]:
    t = sh.text_frame.text
    if t.strip() == "Имя Фамилия":
        put(sh, "[Имя Фамилия]")
    elif t.startswith("Роль в команде"):
        put(sh, ["[Роль в команде]", "[@ник в мессенджере]", "[телефон]", "[место учёбы/работы]"])

s = keep(10)
text(s, 4, 3.6, 42, 5, "О команде", size=16, color=RGBColor(255, 255, 255), bold_first=True)
put(shape(s, "Текст 8"), "[Как собрались, участвовали ли вместе в хакатонах/проектах, интересные факты]")
bodies = [sh for sh in s.shapes if sh.name == "Текст 8" and sh.shape_type == 17]
for sh in bodies:
    t = sh.text_frame.text
    if t.startswith("Что вас вдохновило"):
        put(sh, "[Почему выбрали эту задачу; что вдохновило в проблеме]")
    elif t.startswith("Расскажите о самых"):
        put(sh, ["Нарушений в разметке 6–36 на критерий — нейросеть учить не на чем: перешли к измерениям по ТЗ.",
                 "Тег масштаба Exposed Area оказался неверен в 20% снимков — нашли по круглой головке бедра.",
                 "Разметка врачей местами расходится с порогами ТЗ — показываем это честно, а не подгоняем."], size=11)

s = keep(11)
put(shape(s, "Текст 16"), "Маркетинговая суть решения")
put(shape(s, "Текст 2"), ["Измерения по критериям ТЗ: угол оси позвоночника, отступы области интереса, "
                          "видимость гребней, малый вертел, детектор посторонних предметов.",
                          "Без нейросетей в итоговой системе: SigLIP проверили — измерения без него точнее.",
                          "Честная оценка: кросс-валидация по пациентам, 95% ДИ, перестановочный тест.",
                          "Локальный сервис: CLI, API, веб-интерфейс, DICOM SC/SR, Docker без сети."])
put(shape(s, "Текст 6"), ["Лаборант видит брак сразу — переснять, пока пациент в кабинете.",
                          "Врач получает готовое объяснение и тратит меньше времени на ручной контроль.",
                          "Организация — единый стандарт качества денситометрии и сопоставимые измерения в динамике.",
                          "Встраивание в PACS/ЕРИС через доп. серию и DICOM SR."])

# ---------------------------------------------------------------- решение
s = dup(24)
title(s, "Проблема и решение")
put(shape(s, "Текст 2"), ["Проблема", "Ошибки укладки и области интереса искажают МПК; контроль — вручную, "
                          "знания критериев разные у разных специалистов."], bold_first=True)
put(shape(s, "Текст 3"), ["Решение", "Сервис проверяет каждый снимок по 5 критериям ТЗ и объясняет нарушение "
                          "числом, порогом и разметкой на изображении."], bold_first=True)
put(shape(s, "Текст 4"), ["Результат", f"Итог «есть нарушение»: ROC-AUC {M(ALL, 'auc'):.2f}, F1 {M(ALL, 'f1'):.2f}; "
                          f"< 0.2 с на снимок на CPU, полностью локально."], bold_first=True)

s = dup(17)
title(s, "Данные и что в них нашли")
for nm, v in zip(("Текст 14", "Текст 15", "Текст 16"), ("252", "6–36", "20%")):
    put(shape(s, nm), v, color=PINK, size=26)
put(shape(s, "Текст 1"), "уникальных снимков", color=PLUM)
put(shape(s, "Текст 2"), ["100 исследований, 499 файлов, аппарат GE Lunar Prodigy Advance.",
                          "Позвоночник 99, бедро 153. Копии одного снимка в выгрузке — обрабатываем один раз."], size=11)
put(shape(s, "Текст 3"), "нарушений на критерий", color=PLUM)
put(shape(s, "Текст 4"), ["Ось 10, укладка позвоночника 6, предметы 17, укладка бедра 36, область интереса 7.",
                          "Отсюда выбор: измерения по ТЗ вместо обучения сети."], size=11)
put(shape(s, "Текст 5"), "снимков с неверным масштабом", color=PLUM)
put(shape(s, "Текст 6"), ["Тег Exposed Area даёт пиксель 1.9 мм вместо 0.6 — отбраковываем.",
                          "Метаданные (копии, «врущий» тег) связаны с браком, но это утечка — в модель не берём."], size=11)

s = dup(13)
title(s, "Критерии ТЗ и как мы их проверяем", "Заголовок 1")
drop(s, "Объект 2")
rows = [("Область", "Нарушение", "Как проверяем", "AUC / F1"),
        ("Позвоночник", "Не выравнена ось", "угол линии через крайние позвонки > 5° (рис. 2 ТЗ)", "spine: Не выравнена ось позвоночника"),
        ("Позвоночник", "Некорректная укладка", "не видны гребни подвздошных костей", "spine: Некорректная укладка"),
        ("Позвоночник", "Посторонние предметы", "детектор тонких ярких линий вне позвоночника", "spine: Присутствуют посторонние предметы"),
        ("Бедро", "Некорректная укладка", "малый вертел (U-образно) + офсет головки; седалищная кость в кадре", "femur: Некорректная укладка"),
        ("Бедро", "Область интереса", "< 3 см ниже малого вертела или скан короче 13 см", "femur: Некорректная область интереса")]
tbl = s.shapes.add_table(len(rows), 4, Emu(6 * U), Emu(17 * U), Emu(117 * U), Emu(40 * U)).table
for c, w in enumerate((18, 26, 55, 18)):
    tbl.columns[c].width = Emu(w * U)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        if r and c == 3:
            val = f"{M(val, 'auc'):.2f} / {M(val, 'f1'):.2f}"
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.name, para.runs[0].font.size = FONT, Pt(12 if r else 12)
        para.runs[0].font.bold = r == 0
        para.runs[0].font.color.rgb = RGBColor(255, 255, 255) if r == 0 else DARK
        cell.fill.solid()
        cell.fill.fore_color.rgb = PLUM if r == 0 else (RGBColor(0xF6, 0xF0, 0xFA) if r % 2 else RGBColor(255, 255, 255))
text(s, 6, 59, 117, 8, "Несколько нарушений на снимке перечисляются все; итог «есть нарушение» — если сработал любой критерий. "
                       "Каждое решение сопровождается объяснением: измеренное значение и порог.", size=12, color=GREY)

s = dup(25)
title(s, "Архитектура решения")
steps = [("Чтение", "DICOM/zip, дедупликация копий, группировка по StudyInstanceUID"),
         ("Область и сторона", "теги и имя файла, иначе — по размеру снимка и анатомии"),
         ("Измерения по ТЗ", "ось, гребни, предметы, отступы, малый вертел, офсет головки"),
         ("Решение", "правила ТЗ + обученные пороги и калибровка вероятностей"),
         ("Вывод", "xlsx/csv (формат ТЗ), PNG, DICOM SC и SR, API, веб-интерфейс")]
for k, (t, b) in enumerate(steps):
    put(shape(s, f"Текст {2 * k + 1}"), t, color=RGBColor(255, 255, 255))
    put(shape(s, f"Текст {2 * k + 2}"), b, size=11, color=RGBColor(255, 255, 255))

s = dup(13)
title(s, "Ось позвоночника: меряем как врач в ТЗ", "Заголовок 1")
drop(s, "Объект 2")
pic(s, ROOT / "out/slides/1_agree.png", 6, 16, 80, 42)
text(s, 88, 17, 36, 42, ["Линия через центры крайних позвонков против вертикали — как на рис. 2 ТЗ.",
                         "Поверка: 594 поворота снимков на известный угол — ошибка 0.19° (допуск ТЗ 3°).",
                         f"AUC {M('spine: Не выравнена ось позвоночника', 'auc'):.2f}. F1 ниже, потому что врачи "
                         "размечали на глаз: 6 из 10 нарушений имеют наклон < 5°, часть «норм» — > 5°.",
                         "Порог не подгоняем: решение следует ТЗ и проверяемо."], size=12)

s = dup(21)
title(s, "Ротация бедра")
pic(s, ROOT / "out/pres/u_shape.png", 3, 13, 63, 52)
fa = "femur: Некорректная укладка"
for k, (hd, body) in enumerate([(f"AUC {M(fa, 'auc'):.2f} · F1 {M(fa, 'f1'):.2f}", "самый частый вид брака: 36 из 150 снимков бедра"),
                                ("Малый вертел + офсет головки", "две независимые оценки; вертел сравнивается с нормой в лог-шкале"),
                                ("SigLIP убран", "третьей оценкой ухудшал: F1 0.63 против 0.70 без него")]):
    put(shape(s, f"Текст {2 * k + 2}"), hd, color=PINK, size=18)
    put(shape(s, f"Текст {2 * k + 3}"), body, size=11)

s = dup(13)
title(s, "Ротация бедра: примеры", "Заголовок 1")
drop(s, "Объект 2")
pic(s, ROOT / "out/pres/rotation_cases.png", 6, 16, 117, 42)
text(s, 6, 59, 117, 8, "ТЗ, рис. 5: при переротации контур гладкий, при недоротации малый вертел слишком большой. "
                       "Сервис пишет в заключении площадь выступа и вероятный тип нарушения.", size=12, color=GREY)

s = dup(13)
title(s, "Посторонние предметы и укладка позвоночника", "Заголовок 1")
drop(s, "Объект 2")
pic(s, ROOT / "out/pres/foreign_cases.png", 6, 16, 78, 42)
fo, sp = "spine: Присутствуют посторонние предметы", "spine: Некорректная укладка"
text(s, 86, 17, 38, 42, [f"Предметы: AUC {M(fo, 'auc'):.2f}, F1 {M(fo, 'f1'):.2f}.",
                         "Тонкие яркие линии вне позвоночника (косточки белья, металл) — рис. 3 ТЗ.",
                         "Разбор ошибок нашёл и убрал 3 дефекта: поля дополнения кадра, края таза, склейку рёбер с тазом. F1 0.51 → 0.76.",
                         f"Укладка: гребни подвздошных костей не видны — AUC {M(sp, 'auc'):.2f}, F1 {M(sp, 'f1'):.2f}."], size=12)

s = dup(13)
title(s, "Область интереса бедра", "Заголовок 1")
drop(s, "Объект 2")
pic(s, ROOT / "out/pres/roi_cases.png", 6, 16, 80, 42)
ro = "femur: Некорректная область интереса"
text(s, 88, 17, 36, 42, ["Рамки GE в выгрузке нет — строим область интереса по анатомии: головка, малый и большой вертел.",
                         "Правила ТЗ: ≥ 3 см ниже малого вертела; скан ≥ 13 см (область ~7 см + 2 × 3 см).",
                         f"AUC {M(ro, 'auc'):.2f}, F1 {M(ro, 'f1'):.2f}: найдено 5 из 7 при 1 ложной тревоге.",
                         "Отступы сверху и сбоку — в отчёте: рамка GE уже анатомии."], size=12)

s = dup(18)
title(s, "Честная оценка: без утечек")
items = [("По пациентам", "снимки одного исследования не делятся между обучением и проверкой"),
         ("Пороги внутри фолдов", "всё обучаемое — только на обучающей части; 5 × 5 разбиений"),
         ("ДИ бутстрэпом", "95% интервалы по исследованиям, а не по снимкам"),
         ("Только то, что помогает", "каждую часть проверяем с ней и без неё: SigLIP ухудшал ротацию (F1 0.70 → 0.63) — убран"),
         ("Поверка измерителя", "поворот и обрезка на известную величину: 0.19° и 1–2% (ТЗ: 3° и 5%)"),
         ("Тот же код", "оценка запускает модель сервиса — evaluate_final.py")]
for k, (t, b) in enumerate(items):
    put(shape(s, f"Текст {k + 14}"), f"{k + 1:02d}", color=RGBColor(255, 255, 255))
    put(shape(s, f"Текст {2 * k + 1}"), t, color=RGBColor(255, 255, 255))
    put(shape(s, f"Текст {2 * k + 2}"), b, size=11, color=RGBColor(255, 255, 255))

s = dup(22)
title(s, "Метрики с 95% ДИ")
pic(s, ROOT / "out/pres/metrics.png", 3, 13, 66, 52)
for k, (hd, body) in enumerate([(f"{M(ALL, 'auc'):.2f}", "ROC-AUC «есть нарушение» (249 снимков)"),
                                (f"{M(ALL, 'f1'):.2f}", "F1 «есть нарушение»"),
                                (f"{QP:.2f}", "ROC-AUC по quality_prob (Разъяснения V2)"),
                                (f"{MACRO:.2f}", "macro-F1 по типам нарушений")]):
    put(shape(s, f"Текст {2 * k + 2}"), hd, color=PINK, size=26)
    put(shape(s, f"Текст {2 * k + 3}"), body, size=12, color=RGBColor(255, 255, 255))

s = dup(15)
title(s, "Что пробовали и почему отказались", "Заголовок 5")
for k, line in enumerate(["SigLIP (замороженный и дообученный в Colab): поверх измерений ротации ухудшает F1 0.70 → 0.63",
                          "CNN-ансамбль (9 сетей): лучше только на области интереса, хуже на 4 критериях; смесь хуже нашего",
                          "MedSAM / SAM для маски бедра: малый вертел «сглаживается» — AUC 0.54–0.59 против 0.76",
                          "2D-синтетика для ротации: ротация — 3D, поворот картинки её не имитирует",
                          "VLM через API запрещена ТЗ (данные наружу); метаданные (копии) — утечка, не используем"]):
    put(shape(s, f"Текст {k + 7}"), line, size=13, color=RGBColor(255, 255, 255))

s = dup(26)
title(s, "Демонстрация")
drop(s, "Рисунок 1")
pic(s, ROOT / "out/pres/ui_4.png", 44.5, 12.3, 42, 31.5)
put(shape(s, "Текст 3"), ["Пакетная обработка", "каталог или zip → results.xlsx/csv в формате ТЗ + quality_prob"], size=11, bold_first=True, color=RGBColor(255, 255, 255))
put(shape(s, "Текст 4"), ["Объяснение врачу", "снимок с разметкой и заключение: значение и порог по каждому критерию"], size=11, bold_first=True, color=RGBColor(255, 255, 255))
put(shape(s, "Текст 5"), ["В PACS", "доп. серия DICOM SC с разметкой и текстовое заключение DICOM SR"], size=11, bold_first=True, color=RGBColor(255, 255, 255))

s = dup(16)
title(s, "Скорость и ограничения")
cards = [("< 0.2 с", "на снимок на CPU; 100 исследований за 21 с при лимите ТЗ 3 мин на одно"),
         ("Офлайн", "Docker 750 МБ, без сети и GPU; CPU 2+ ядра, 2 ГБ ОЗУ"),
         ("Ограничения", "один аппарат; мало нарушений — широкие ДИ; Th12 не проверяем; предметы поверх позвоночника"),
         ("Масштаб", "организатор: 0.6 × 1.05 мм; головка бедра круглая (0.95) — считаем квадратный пиксель")]
for k, (t, b) in enumerate(cards):
    put(shape(s, f"Текст {k + 14}"), f"{k + 1:02d}", color=PINK)
    put(shape(s, f"Текст {2 * k + 1}"), t, color=PLUM)
    put(shape(s, f"Текст {2 * k + 2}"), b, size=11)

s = dup(17)
title(s, "Кому это нужно")
for nm, v in zip(("Текст 14", "Текст 15", "Текст 16"), ("01", "02", "03")):
    put(shape(s, nm), v, color=PINK, size=20)
put(shape(s, "Текст 1"), "Рентгенолаборант", color=PLUM)
put(shape(s, "Текст 2"), ["Сигнал о браке сразу после снимка — переснять, пока пациент в кабинете.",
                          "Меньше повторных визитов и дополнительной лучевой нагрузки."], size=11)
put(shape(s, "Текст 3"), "Врач-рентгенолог", color=PLUM)
put(shape(s, "Текст 4"), ["Готовое объяснение: что нарушено и на сколько.",
                          "Меньше ручного контроля; сопоставимые измерения МПК в динамике."], size=11)
put(shape(s, "Текст 5"), "Медорганизация / ДЗМ", color=PLUM)
put(shape(s, "Текст 6"), ["Единый стандарт качества денситометрии по городу.",
                          "Статистика брака по кабинетам — адресное обучение персонала. [оценка экономики — при наличии данных]"], size=11)

s = dup(25)
title(s, "План развития и внедрения")
for k, (t, b) in enumerate([("Пилот", "«вторым читателем» рядом с лаборантом, сбор обратной связи"),
                            ("Th12", "подсчёт позвонков от крестца — верхняя граница укладки"),
                            ("Аппараты", "Hologic и др.: перенастройка порогов через train.py"),
                            ("Интеграция", "PACS/ЕРИС: доп. серия и DICOM SR в поток"),
                            ("Разметка", "накопление подтверждённых врачом случаев → дообучение")]):
    put(shape(s, f"Текст {2 * k + 1}"), t, color=RGBColor(255, 255, 255))
    put(shape(s, f"Текст {2 * k + 2}"), b, size=11, color=RGBColor(255, 255, 255))

s = dup(7)
put(shape(s, "Заголовок 2"), "Спасибо!", color=RGBColor(255, 255, 255))
put(shape(s, "Текст 4"), "[Название команды] · [контакт капитана] · github.com/qkwe/bone-densitometry (ветка v2)", color=RGBColor(255, 255, 255))

# ---------------------------------------------------------------- порядок и удаление служебных слайдов
sld_ids = prs.slides._sldIdLst
by_part = {id(sl.part): sl for sl in prs.slides}
id_of = {}
for sid in list(sld_ids):
    part = prs.part.related_part(sid.rId)
    id_of[id(part)] = sid
keep_ids = [id_of[id(sl.part)] for sl in final]
keep_rids = {sid.rId for sid in keep_ids}
for sid in list(sld_ids):
    sld_ids.remove(sid)
    if sid.rId not in keep_rids:  # служебные слайды шаблона удаляются из файла целиком
        prs.part.drop_rel(sid.rId)
for sid in keep_ids:
    sld_ids.append(sid)
prs.save(OUT)
print("слайдов:", len(keep_ids), "→", OUT)
